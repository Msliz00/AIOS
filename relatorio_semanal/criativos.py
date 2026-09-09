"""Leitura da planilha de criativos (uma aba por semana) + match tolerante da aba.

Header na linha 1, dados a partir da linha 2. O nome da aba não tem padrão fixo
(`S35 · 24–3008`, `2807 a 0308`, `S27 · 2906–0507 (Jul)`), então o match é por
número ISO da semana OU pelo range de datas embutido no título.
"""
from __future__ import annotations

import datetime as dt
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

# sinônimos de header → campo interno (comparados após normalização)
HEADERS = {
    "nome": ("nome", "linkcriativo", "criativo"),
    "investimento": ("investimento", "gasto", "invest"),
    "leads": ("leads", "lead"),
    "registros": ("registro", "registros", "reg"),
    "ftd": ("ftd",),
    "thumbstop": ("thumbstop",),
    "holdrate": ("holdrate",),
    "ctr": ("ctr",),
    "novo": ("novo",),
}


def normalizar(texto: str) -> str:
    """Colapsa espaços, remove acento/traço/·/aspas e baixa a caixa."""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"[·•\-–—_'\"“”‘’()\[\]]", " ", t)
    return re.sub(r"\s+", " ", t).strip().lower()


@dataclass
class Criativo:
    nome: str
    investimento: float
    leads: float
    registros: float
    ftd: float
    thumbstop: float | None = None
    holdrate: float | None = None
    ctr: float | None = None
    novo: bool = False

    @property
    def c_lead(self) -> float | None:
        return self.investimento / self.leads if self.leads else None

    @property
    def c_reg(self) -> float | None:
        return self.investimento / self.registros if self.registros else None

    @property
    def c_ftd(self) -> float | None:
        return self.investimento / self.ftd if self.ftd else None


@dataclass
class AbaCriativos:
    titulo: str
    criativos: list[Criativo] = field(default_factory=list)
    motivo_match: str = ""

    @property
    def investimento(self) -> float:
        return sum(c.investimento for c in self.criativos)

    @property
    def registros(self) -> float:
        return sum(c.registros for c in self.criativos)

    @property
    def ftd(self) -> float:
        return sum(c.ftd for c in self.criativos)


# --------------------------------------------------------------------------
# match tolerante da aba
# --------------------------------------------------------------------------
MESES_ABREV = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
_RX_SEMANA = re.compile(r"\bs\s?(\d{1,2})\b")
_RX_RANGE = re.compile(r"(\d{1,2})(\d{2})?\s*(?:a|ate|to)?\s+(\d{2})(\d{2})\b")


def _range_do_titulo(titulo: str, ano: int) -> tuple[dt.date, dt.date] | None:
    """Extrai `DDMM–DDMM`, `DD–DDMM` ou `DDMM a DDMM` de um título normalizado."""
    m = _RX_RANGE.search(titulo)
    if not m:
        return None
    d1, m1, d2, m2 = m.groups()
    try:
        fim = dt.date(ano, int(m2), int(d2))
        ini = dt.date(ano, int(m1) if m1 else int(m2), int(d1))
        if ini > fim:  # `2912 a 0401`: virou o ano
            ini = ini.replace(year=ano - 1)
        return ini, fim
    except ValueError:
        return None


def _sobreposicao(a: tuple[dt.date, dt.date], b: tuple[dt.date, dt.date]) -> int:
    ini, fim = max(a[0], b[0]), min(a[1], b[1])
    return max(0, (fim - ini).days + 1)


def achar_aba(nomes: list[str], inicio: dt.date, fim: dt.date) -> tuple[str, str] | None:
    """Devolve (nome_da_aba, motivo) ou None. Nunca compara string exata."""
    semana = inicio.isocalendar()[1]
    janela = (inicio, fim)
    cand_semana, cand_range = [], []
    for nome in nomes:
        n = normalizar(nome)
        ms = _RX_SEMANA.search(n)
        if ms and int(ms.group(1)) == semana:
            cand_semana.append(nome)
        rng = _range_do_titulo(n, fim.year)
        if rng:
            ov = _sobreposicao(rng, janela)
            if ov:
                cand_range.append((ov, -abs((rng[1] - rng[0]).days - (fim - inicio).days), nome))
    if cand_semana:
        # vários "S27 …": preferir o que não é "sem filtro" e cujo range casa melhor
        mes_fim = MESES_ABREV[fim.month - 1]

        def score(nome):
            n = normalizar(nome)
            rng = _range_do_titulo(n, fim.year)
            return (
                "sem filtro" not in n,
                _sobreposicao(rng, janela) if rng else 0,
                mes_fim in n,  # "S27 … (Jul)" vs "S27 … (Jun)": fica com o mês em que a semana termina
            )
        melhor = max(cand_semana, key=score)
        return melhor, f"semana ISO S{semana}"
    if cand_range:
        cand_range.sort(reverse=True)
        return cand_range[0][2], "range de datas no título"
    return None


# --------------------------------------------------------------------------
# leitura
# --------------------------------------------------------------------------
def _num(v) -> float:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return 0.0


def _opt(v) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _mapear_header(row) -> dict[str, int]:
    mapa: dict[str, int] = {}
    for idx, cel in enumerate(row):
        if cel is None:
            continue
        chave = normalizar(cel).replace(" ", "").replace("/", "")
        for campo, sinonimos in HEADERS.items():
            if campo in mapa:
                continue
            if chave in sinonimos or any(chave.startswith(s) for s in sinonimos if len(s) > 3):
                mapa[campo] = idx
                break
    return mapa


def ler_aba(caminho: Path, titulo: str) -> AbaCriativos:
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = wb[titulo]
    linhas = ws.iter_rows(values_only=True)
    header = next(linhas, None)
    mapa = _mapear_header(header or ())
    if "nome" not in mapa or "investimento" not in mapa:
        wb.close()
        raise ValueError(f"aba '{titulo}': header sem Nome/Investimento ({header})")
    aba = AbaCriativos(titulo=titulo)

    def col(row, campo):
        i = mapa.get(campo)
        return row[i] if i is not None and i < len(row) else None

    for row in linhas:
        nome = col(row, "nome")
        if not isinstance(nome, str) or not nome.strip():
            continue
        if normalizar(nome) in ("total", "totais"):
            continue
        aba.criativos.append(
            Criativo(
                nome=nome.strip(),
                investimento=_num(col(row, "investimento")),
                leads=_num(col(row, "leads")),
                registros=_num(col(row, "registros")),
                ftd=_num(col(row, "ftd")),
                thumbstop=_opt(col(row, "thumbstop")),
                holdrate=_opt(col(row, "holdrate")),
                ctr=_opt(col(row, "ctr")),
                novo=bool(col(row, "novo")),
            )
        )
    wb.close()
    aba.criativos.sort(key=lambda c: c.investimento, reverse=True)
    return aba


def ler_semana(caminho: Path, inicio: dt.date, fim: dt.date) -> AbaCriativos | None:
    wb = openpyxl.load_workbook(caminho, read_only=True)
    nomes = list(wb.sheetnames)
    wb.close()
    achado = achar_aba(nomes, inicio, fim)
    if not achado:
        return None
    titulo, motivo = achado
    aba = ler_aba(caminho, titulo)
    aba.motivo_match = motivo
    return aba
