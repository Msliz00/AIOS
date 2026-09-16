"""Cruza as mídias de uma pasta do Drive com uma aba semanal da planilha de criativos.

Uso:
    python -m relatorio_semanal.cruzar_pasta pasta.json criativos.xlsx "S37" saida.xlsx

`pasta.json` = lista de arquivos da pasta ({id, title, fileSize?, modifiedTime?}),
como devolvido pelo Drive (API `files.list` ou conector MCP `search_files`).
Match em duas chaves: id do Drive dentro da coluna LINK, senão nome normalizado
sem extensão. Saída: XLSX com as mídias da pasta presentes na aba (com métricas),
as da pasta fora da aba e as da aba fora da pasta.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .criativos import achar_aba, ler_aba, normalizar

_RX_ID = re.compile(r"/d/([A-Za-z0-9_-]{20,})")
FONTE = Font(name="Arial", size=10)
FONTE_H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor="FF6B00")


def _div(a: float, b: float):
    return round(a / b, 2) if b else None


def _chave(nome: str) -> str:
    nome = re.sub(r"\.(mp4|mov|m4v|avi)$", "", nome.strip(), flags=re.I)
    return normalizar(nome).replace(" ", "")


def cruzar(pasta: list[dict], criativos_xlsx: Path, aba_titulo: str, saida: Path) -> dict:
    wb_src = openpyxl.load_workbook(criativos_xlsx, read_only=True)
    nomes = wb_src.sheetnames
    wb_src.close()
    titulo = next((n for n in nomes if normalizar(n).startswith(normalizar(aba_titulo))), None)
    if titulo is None:
        raise KeyError(f"aba '{aba_titulo}' não encontrada; abas: {nomes[-6:]}")
    aba = ler_aba(criativos_xlsx, titulo)

    # links crus por linha (ler_aba não guarda o LINK)
    wb_src = openpyxl.load_workbook(criativos_xlsx, read_only=True, data_only=True)
    ws = wb_src[titulo]
    links: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and isinstance(row[0], str) and len(row) > 2 and isinstance(row[2], str):
            links[row[0].strip()] = row[2]
    wb_src.close()

    por_id = {}
    for c in aba.criativos:
        m = _RX_ID.search(links.get(c.nome, "") or "")
        if m:
            por_id[m.group(1)] = c
    por_nome = {_chave(c.nome): c for c in aba.criativos}

    casados, so_pasta = [], []
    usados = set()
    for f in pasta:
        c = por_id.get(f["id"])
        metodo = "id do Drive no LINK"
        if c is None:
            c = por_nome.get(_chave(f["title"]))
            metodo = "nome"
        if c is None:
            so_pasta.append(f)
        else:
            casados.append((f, c, metodo))
            usados.add(c.nome)
    so_aba = [c for c in aba.criativos if c.nome not in usados]
    casados.sort(key=lambda t: t[1].investimento, reverse=True)

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Na pasta e na aba"
    cab = ["Arquivo na pasta", "Nome na aba", "Match", "Investimento", "Leads", "C/Lead",
           "Registro", "C/Reg", "FTD", "C/FTD", "ThumbStop", "HoldRate", "CTR", "Link Drive", "Tamanho (MB)"]
    ws1.append(cab)
    for f, c, metodo in casados:
        ws1.append([
            f["title"], c.nome, metodo, round(c.investimento, 2), c.leads, _div(c.investimento, c.leads),
            c.registros, _div(c.investimento, c.registros), c.ftd, _div(c.investimento, c.ftd),
            c.thumbstop, c.holdrate, c.ctr,
            f"https://drive.google.com/file/d/{f['id']}/view",
            round(int(f.get("fileSize", 0) or 0) / 1e6, 1),
        ])
    n = len(casados)
    tot = n + 2
    inv = sum(c.investimento for _, c, _ in casados)
    leads = sum(c.leads for _, c, _ in casados)
    regs = sum(c.registros for _, c, _ in casados)
    ftds = sum(c.ftd for _, c, _ in casados)
    ws1[f"A{tot}"] = "TOTAL"
    ws1[f"D{tot}"], ws1[f"E{tot}"], ws1[f"F{tot}"] = round(inv, 2), leads, _div(inv, leads)
    ws1[f"G{tot}"], ws1[f"H{tot}"] = regs, _div(inv, regs)
    ws1[f"I{tot}"], ws1[f"J{tot}"] = ftds, _div(inv, ftds)
    ws1[f"A{tot+2}"] = (
        f"Fonte: aba '{titulo}' da planilha de criativos × listagem da pasta do Drive. "
        "Match por id do Drive na coluna LINK; quando a aba não tem link, por nome sem extensão. "
        "C/Lead, C/Reg e C/FTD recalculados (Investimento ÷ métrica). Valores gravados, não fórmulas."
    )
    ws1[f"A{tot+2}"].font = Font(name="Arial", size=9, italic=True)

    ws2 = wb.create_sheet("Na pasta, fora da aba")
    ws2.append(["Arquivo na pasta", "Link Drive", "Tamanho (MB)", "Modificado"])
    for f in so_pasta:
        ws2.append([f["title"], f"https://drive.google.com/file/d/{f['id']}/view",
                    round(int(f.get("fileSize", 0) or 0) / 1e6, 1), f.get("modifiedTime", "")])

    ws3 = wb.create_sheet("Na aba, fora da pasta")
    ws3.append(["Nome na aba", "Investimento", "Leads", "Registro", "FTD"])
    for c in so_aba:
        ws3.append([c.nome, round(c.investimento, 2), c.leads, c.registros, c.ftd])
    m = len(so_aba) + 2
    ws3[f"A{m}"] = "TOTAL"
    ws3[f"B{m}"] = round(sum(c.investimento for c in so_aba), 2)
    ws3[f"C{m}"] = sum(c.leads for c in so_aba)
    ws3[f"D{m}"] = sum(c.registros for c in so_aba)
    ws3[f"E{m}"] = sum(c.ftd for c in so_aba)

    ws4 = wb.create_sheet("Resumo")
    ws4.append(["Métrica", "Valor"])
    ws4.append(["Aba analisada", titulo])
    ws4.append(["Arquivos na pasta", len(pasta)])
    ws4.append(["Pasta ∩ aba", n])
    ws4.append(["Pasta fora da aba", len(so_pasta)])
    ws4.append(["Aba fora da pasta", len(so_aba)])
    ws4.append(["Invest. da aba (total)", round(aba.investimento, 2)])
    ws4.append(["Invest. das mídias da pasta", round(inv, 2)])
    ws4.append(["Share da pasta no invest. da aba", _div(inv, aba.investimento)])
    ws4.append(["FTD da aba (total)", aba.ftd])
    ws4.append(["FTD das mídias da pasta", ftds])
    ws4["B9"].number_format = "0.0%"

    for w in (ws1, ws2, ws3, ws4):
        for cell in w[1]:
            cell.font, cell.fill = FONTE_H, FILL_H
        for row in w.iter_rows(min_row=2):
            for cell in row:
                if cell.font.name != "Arial" or not cell.font.italic:
                    cell.font = FONTE if not cell.font.bold else cell.font
        for i, col in enumerate(w.columns, 1):
            largura = max((len(str(c.value)) for c in col if c.value is not None and not str(c.value).startswith("=")), default=8)
            w.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 60)
        w.freeze_panes = "A2"
    for r in range(2, tot + 1):
        for col in "DFHJ":
            ws1[f"{col}{r}"].number_format = '"R$" #,##0.00'
    for r in range(2, m + 1):
        ws3[f"B{r}"].number_format = '"R$" #,##0.00'
    ws1[f"A{tot}"].font = Font(name="Arial", size=10, bold=True)
    ws3[f"A{m}"].font = Font(name="Arial", size=10, bold=True)
    wb.save(saida)
    return {"aba": titulo, "casados": n, "so_pasta": len(so_pasta), "so_aba": len(so_aba),
            "invest_pasta": sum(c.investimento for _, c, _ in casados), "invest_aba": aba.investimento,
            "ftd_pasta": sum(c.ftd for _, c, _ in casados), "ftd_aba": aba.ftd,
            "lista": [(f["title"], c.nome, metodo, c.investimento, c.ftd) for f, c, metodo in casados],
            "so_pasta_nomes": [f["title"] for f in so_pasta]}


def main(argv=None):
    a = argv or sys.argv[1:]
    if len(a) != 4:
        print(__doc__)
        return 2
    pasta = json.loads(Path(a[0]).read_text())
    if isinstance(pasta, dict):
        pasta = pasta.get("files", [])
    r = cruzar(pasta, Path(a[1]), a[2], Path(a[3]))
    print(json.dumps({k: v for k, v in r.items() if k not in ("lista",)}, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
