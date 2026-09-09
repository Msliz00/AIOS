"""Leitura da planilha de operação (diário por expert).

Layout: uma aba por mês em PT (JANEIRO…DEZEMBRO), header na linha 2, dados a
partir da linha 3, colunas por índice fixo. A janela pode cruzar meses.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

MESES_PT = [
    "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
    "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO",
]
COL_DATA, COL_GASTO, COL_LEADS, COL_REG, COL_FTD, COL_DEP, COL_NGR = 0, 1, 2, 5, 7, 10, 11
LINHA_PRIMEIRO_DADO = 3


def _num(v) -> float:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return 0.0  # None, '#DIV/0!', texto → 0


@dataclass
class Dia:
    data: dt.date
    gasto: float
    leads: float
    registros: float
    ftd: float
    deposito: float
    ngr: float


@dataclass
class Janela:
    inicio: dt.date
    fim: dt.date
    dias: list[Dia] = field(default_factory=list)
    abas_lidas: list[str] = field(default_factory=list)
    abas_faltantes: list[str] = field(default_factory=list)

    # --- agregados -------------------------------------------------------
    @property
    def gasto(self) -> float:
        return sum(d.gasto for d in self.dias)

    @property
    def leads(self) -> float:
        return sum(d.leads for d in self.dias)

    @property
    def registros(self) -> float:
        return sum(d.registros for d in self.dias)

    @property
    def ftd(self) -> float:
        return sum(d.ftd for d in self.dias)

    @property
    def deposito(self) -> float:
        return sum(d.deposito for d in self.dias)

    @property
    def ngr(self) -> float:
        return sum(d.ngr for d in self.dias)

    @property
    def c_lead(self) -> float | None:
        return self.gasto / self.leads if self.leads else None

    @property
    def c_reg(self) -> float | None:
        return self.gasto / self.registros if self.registros else None

    @property
    def c_ftd(self) -> float | None:
        return self.gasto / self.ftd if self.ftd else None

    @property
    def pct_reg_ftd(self) -> float | None:
        return self.ftd / self.registros if self.registros else None

    @property
    def ngr_menos_gasto(self) -> float:
        return self.ngr - self.gasto

    @property
    def dias_com_gasto(self) -> int:
        return sum(1 for d in self.dias if d.gasto > 0)

    @property
    def dias_esperados(self) -> int:
        return (self.fim - self.inicio).days + 1

    @property
    def sem_midia(self) -> bool:
        return self.gasto <= 0


def _meses_da_janela(inicio: dt.date, fim: dt.date) -> list[tuple[int, int]]:
    meses, cursor = [], dt.date(inicio.year, inicio.month, 1)
    while cursor <= fim:
        meses.append((cursor.year, cursor.month))
        cursor = dt.date(cursor.year + (cursor.month == 12), cursor.month % 12 + 1, 1)
    return meses


def _achar_aba(wb, mes: int):
    alvo = MESES_PT[mes - 1]
    for nome in wb.sheetnames:
        if nome.strip().upper() == alvo:
            return wb[nome]
    return None


def ler_janela(caminho: Path, inicio: dt.date, fim: dt.date) -> Janela:
    """Soma o diário em [inicio, fim] (inclusivo), cruzando abas de mês."""
    if fim < inicio:
        raise ValueError("data_fim antes de data_inicio")
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    janela = Janela(inicio, fim)
    for ano, mes in _meses_da_janela(inicio, fim):
        ws = _achar_aba(wb, mes)
        if ws is None:
            janela.abas_faltantes.append(MESES_PT[mes - 1])
            continue
        janela.abas_lidas.append(ws.title)
        for row in ws.iter_rows(min_row=LINHA_PRIMEIRO_DADO, values_only=True):
            data = row[COL_DATA] if row else None
            if isinstance(data, dt.datetime):
                data = data.date()
            if not isinstance(data, dt.date):
                continue  # linha TOTAL, vazia ou texto
            if data.year != ano or data.month != mes or not (inicio <= data <= fim):
                continue
            janela.dias.append(
                Dia(
                    data=data,
                    gasto=_num(row[COL_GASTO]),
                    leads=_num(row[COL_LEADS]),
                    registros=_num(row[COL_REG]),
                    ftd=_num(row[COL_FTD]),
                    deposito=_num(row[COL_DEP]),
                    ngr=_num(row[COL_NGR]),
                )
            )
    wb.close()
    janela.dias.sort(key=lambda d: d.data)
    return janela
