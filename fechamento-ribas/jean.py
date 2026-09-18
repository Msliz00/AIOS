#!/usr/bin/env python3
"""Gerador JEAN VLOGS (38579) - all-time diario desde 13/05, SOMENTE metricas (10 colunas).

Uso:
  python3 jean.py --input export_affiliate_38579.csv [--ate 2026-09-17] [--desde 2026-05-13] [--out DIR]

Regras de formato (fixas):
  - 10 colunas A->J, SEM 'CUSTO P/ FTD', SEM CPA/RevShare/Comissao Dia.
  - Bloco METRICAS GERAIS SEM linha de comissao.
  - Visual Play: fundo preto, header/TOTAL dourado, NGR verde/vermelho, TOTAL rotulado '2026'.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook

from common import (BG, CENTER, FILL_BG, FILL_GOLD, INT, LEFT, MONEY, PCT,
                    aplicar_estilo, checar_ultimo_dia, f, filtrar, read_export)

AFFILIATE = "38579"
DESDE_PADRAO = "2026-05-13"
HEADERS = ["DIA", "REGISTROS", "FTD", "% REG P/ FTD", "VALOR DEP.", "NGR",
           "MÉDIA DEP/DIA", "MÉDIA NGR/DIA", "% DEP → NGR", "TICKET MÉD. DEP."]


def build(df, desde, ate, out_dir: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "JEAN VLOGS"

    titulo = (f"FECHAMENTO JEAN VLOGS ({AFFILIATE}) — "
              f"{desde:%d/%m/%y} ATÉ {ate:%d/%m/%y}")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    c = ws.cell(row=1, column=1, value=titulo)
    c.font = f(size=14, bold=True, color="F4C430")
    c.alignment = CENTER

    HEADER_ROW = 3
    for i, h in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=i, value=h)

    first = HEADER_ROW + 1
    for i, rec in enumerate(df.itertuples(index=False)):
        r = first + i
        ws.cell(row=r, column=1, value=rec.day.strftime("%d/%m/%Y"))
        ws.cell(row=r, column=2, value=int(round(rec.registros))).number_format = INT
        ws.cell(row=r, column=3, value=int(round(rec.ftd))).number_format = INT
        ws.cell(row=r, column=4, value=f"=IF(B{r}=0,0,C{r}/B{r})").number_format = PCT
        ws.cell(row=r, column=5, value=float(rec.deposits)).number_format = MONEY
        ws.cell(row=r, column=6, value=float(rec.ngr)).number_format = MONEY
        ws.cell(row=r, column=7, value=f"=E{r}").number_format = MONEY
        ws.cell(row=r, column=8, value=f"=F{r}").number_format = MONEY
        ws.cell(row=r, column=9, value=f"=IF(E{r}=0,0,F{r}/E{r})").number_format = PCT
        ws.cell(row=r, column=10, value=f"=IF(C{r}=0,0,E{r}/C{r})").number_format = MONEY

    last = first + len(df) - 1
    t = last + 1
    ws.cell(row=t, column=1, value="2026")
    ws.cell(row=t, column=2, value=f"=SUM(B{first}:B{last})").number_format = INT
    ws.cell(row=t, column=3, value=f"=SUM(C{first}:C{last})").number_format = INT
    ws.cell(row=t, column=4, value=f"=IF(B{t}=0,0,C{t}/B{t})").number_format = PCT
    ws.cell(row=t, column=5, value=f"=SUM(E{first}:E{last})").number_format = MONEY
    ws.cell(row=t, column=6, value=f"=SUM(F{first}:F{last})").number_format = MONEY
    ws.cell(row=t, column=7, value=f"=AVERAGE(E{first}:E{last})").number_format = MONEY
    ws.cell(row=t, column=8, value=f"=AVERAGE(F{first}:F{last})").number_format = MONEY
    ws.cell(row=t, column=9, value=f"=IF(E{t}=0,0,F{t}/E{t})").number_format = PCT
    ws.cell(row=t, column=10, value=f"=IF(C{t}=0,0,E{t}/C{t})").number_format = MONEY

    # ---- bloco METRICAS GERAIS (sem comissao)
    m = t + 2
    ws.merge_cells(start_row=m, start_column=1, end_row=m, end_column=2)
    cm = ws.cell(row=m, column=1, value="MÉTRICAS GERAIS")
    cm.font = f(size=12, bold=True, color="F4C430")
    cm.alignment = LEFT

    metricas = [
        ("Total Dias", len(df), INT),
        ("Total Registros", f"=B{t}", INT),
        ("Total FTD", f"=C{t}", INT),
        ("Total Valor Depositado", f"=E{t}", MONEY),
        ("Total NGR", f"=F{t}", MONEY),
        ("Ticket Médio Depósito", f"=J{t}", MONEY),
        ("Ticket Médio NGR", f"=IF(C{t}=0,0,F{t}/C{t})", MONEY),
        ("Conversão Registro→FTD", f"=D{t}", PCT),
        ("% Depósito→NGR", f"=I{t}", PCT),
    ]
    for i, (label, val, fmt) in enumerate(metricas, start=1):
        r = m + i
        lc = ws.cell(row=r, column=1, value=label)
        lc.alignment = LEFT
        vc = ws.cell(row=r, column=2, value=val)
        vc.number_format = fmt

    aplicar_estilo(ws, 10, HEADER_ROW, first, t, ngr_col=6, ultima_linha=m + len(metricas))
    ws.cell(row=1, column=1).fill = FILL_BG
    for i in range(1, len(metricas) + 1):
        ws.cell(row=m + i, column=1).alignment = LEFT
    ws.cell(row=m, column=1).alignment = LEFT

    out = out_dir / f"FECHAMENTO_JEAN_ATE_{ate:%d%m%y}.xlsx"
    wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="export Affiliate (Affiliate=38579), diario")
    ap.add_argument("--desde", default=DESDE_PADRAO)
    ap.add_argument("--ate", default=None, help="opcional; sem isso usa o ultimo dia do export")
    ap.add_argument("--out", default=".")
    a = ap.parse_args()

    df = read_export(a.input, expect_affiliate=AFFILIATE)
    ate = checar_ultimo_dia(df, a.ate)
    df = filtrar(df, a.desde, ate)
    desde = df["day"].min()

    out_dir = Path(a.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = build(df, desde, ate, out_dir)

    print(f"OK  {path}")
    print(f"    periodo   : {desde:%d/%m/%Y} -> {ate:%d/%m/%Y}  ({len(df)} dias)")
    print(f"    registros : {df.registros.sum():,.0f}")
    print(f"    ftd       : {df.ftd.sum():,.0f}")
    print(f"    depositos : R$ {df.deposits.sum():,.2f}")
    print(f"    ngr       : R$ {df.ngr.sum():,.2f}")
    print("    -> rodar: python3 /mnt/skills/public/xlsx/scripts/recalc.py "
          f"{path} 90   (esperar total_errors: 0)")


if __name__ == "__main__":
    main()
