#!/usr/bin/env python3
"""Fechamento SEMANAL BINGO/REALS (pendente) - arquivo acumulado.

Uso:
  python3 pendente.py --bingo exp_bingo.csv --reals exp_reals.csv \
      --inicio 2026-08-17 --fim 2026-08-23 [--arquivo FECHAMENTO_FINAL_PENDENTE_...xlsx]

Semana = SEGUNDA a DOMINGO. Cria uma aba 'B <semana>' e 'R <semana>' por semana
e reconstroi a aba INDICE (renumera # e corrige os 3 SUM).

Regras:
  BINGO: E=ROUND(D*40,2)  H=ROUND(F*G,2)  I=IF(D=0,0,IF(E+H<0,0,E+H))     G=0.35
  REALS: E=ROUND(D*30,2)  H=ROUND(F*G,2)  I=E+H  (mantem negativo)        G=0.35 (3 Daniel) / 0.25
"""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from common import (CENTER, FILL_BG, FILL_GOLD, INT, LEFT, MONEY,
                    aplicar_estilo, f, read_contas)
from roster import BINGO, EXCLUIDOS, OBS, REALS, REVSHARE35_REALS

HEADERS = ["USERNAME", "BTAG", "PERÍODO", "FTD", "COMISSÃO FTD", "NGR",
           "%", "COMISSÃO NGR", "FINAL", "OBS"]
HEADER_ROW = 2
FIRST = 3


def _semana_label(inicio, fim):
    return f"{inicio:%d}a{fim:%d}.{fim:%m}"


def _monta_aba(wb, nome, contas, dados, periodo, cpa, revshare_fn, zera):
    if nome in wb.sheetnames:
        del wb[nome]
    ws = wb.create_sheet(nome)
    for i, h in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=i, value=h)

    for i, (btag, username) in enumerate(contas):
        r = FIRST + i
        rec = dados.get(btag) or dados.get(username.lower()) or {"ftd": 0.0, "ngr": 0.0}
        rs = revshare_fn(btag)
        ws.cell(row=r, column=1, value=username).alignment = LEFT
        ws.cell(row=r, column=2, value=btag)
        ws.cell(row=r, column=3, value=periodo)
        ws.cell(row=r, column=4, value=int(round(rec["ftd"]))).number_format = INT
        ws.cell(row=r, column=5, value=f"=ROUND(D{r}*{cpa},2)").number_format = MONEY
        ws.cell(row=r, column=6, value=float(rec["ngr"])).number_format = MONEY
        ws.cell(row=r, column=7, value=rs).number_format = "0%"
        ws.cell(row=r, column=8, value=f"=ROUND(F{r}*G{r},2)").number_format = MONEY
        final = (f"=IF(D{r}=0,0,IF(E{r}+H{r}<0,0,E{r}+H{r}))" if zera else f"=E{r}+H{r}")
        ws.cell(row=r, column=9, value=final).number_format = MONEY
        ws.cell(row=r, column=10, value=OBS.get(btag, "")).alignment = LEFT

    last = FIRST + len(contas) - 1
    t = last + 1
    ws.cell(row=t, column=1, value="TOTAL")
    ws.cell(row=t, column=3, value=periodo)
    ws.cell(row=t, column=4, value=f"=SUM(D{FIRST}:D{last})").number_format = INT
    ws.cell(row=t, column=5, value=f"=SUM(E{FIRST}:E{last})").number_format = MONEY
    ws.cell(row=t, column=6, value=f"=SUM(F{FIRST}:F{last})").number_format = MONEY
    ws.cell(row=t, column=8, value=f"=SUM(H{FIRST}:H{last})").number_format = MONEY
    ws.cell(row=t, column=9, value=f"=SUM(I{FIRST}:I{last})").number_format = MONEY

    aplicar_estilo(ws, 10, HEADER_ROW, FIRST, t, ngr_col=6, ultima_linha=t)
    for r in range(FIRST, t + 1):
        ws.cell(row=r, column=1).alignment = LEFT
        ws.cell(row=r, column=10).alignment = LEFT
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["J"].width = 28
    return t


def _reconstroi_indice(wb):
    if "ÍNDICE" in wb.sheetnames:
        del wb["ÍNDICE"]
    ws = wb.create_sheet("ÍNDICE", 0)
    semanas = []
    for nome in wb.sheetnames:
        if nome.startswith("B ") and f"R {nome[2:]}" in wb.sheetnames:
            semanas.append(nome[2:])
    semanas.sort(key=lambda s: (s.split(".")[1], s.split("a")[0]))

    ws.cell(row=1, column=1, value="ÍNDICE — FECHAMENTO FINAL PENDENTE").font = f(
        size=14, bold=True, color="F4C430")
    for i, h in enumerate(["#", "PLATAFORMA", "SEMANA", "TOTAL FINAL"], start=1):
        ws.cell(row=3, column=i, value=h)

    r = 4
    b_rows, r_rows = [], []
    for n, sem in enumerate(semanas, start=1):
        ws.cell(row=r, column=1, value=n).number_format = INT
        ws.cell(row=r, column=2, value="BINGO")
        ws.cell(row=r, column=3, value=sem)
        ws.cell(row=r, column=4, value=f"='B {sem}'!I12").number_format = MONEY
        b_rows.append(r)
        r += 1
        ws.cell(row=r, column=1, value=n).number_format = INT
        ws.cell(row=r, column=2, value="REALS")
        ws.cell(row=r, column=3, value=sem)
        ws.cell(row=r, column=4, value=f"='R {sem}'!I56").number_format = MONEY
        r_rows.append(r)
        r += 1

    r += 1
    def soma(rows):
        return "=" + "+".join(f"D{x}" for x in rows) if rows else "=0"
    ws.cell(row=r, column=2, value="TOTAL BINGO")
    ws.cell(row=r, column=4, value=soma(b_rows)).number_format = MONEY
    ws.cell(row=r + 1, column=2, value="TOTAL REALS")
    ws.cell(row=r + 1, column=4, value=soma(r_rows)).number_format = MONEY
    ws.cell(row=r + 2, column=2, value="TOTAL GERAL PENDENTE")
    ws.cell(row=r + 2, column=4, value=f"=D{r}+D{r + 1}").number_format = MONEY

    aplicar_estilo(ws, 4, 3, 4, r + 2, ultima_linha=r + 2)
    for rr in (r, r + 1, r + 2):
        for c in range(1, 5):
            ws.cell(row=rr, column=c).fill = FILL_GOLD
            ws.cell(row=rr, column=c).font = f(bold=True, color="000000")
    ws.cell(row=1, column=1).fill = FILL_BG
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["D"].width = 20
    return semanas


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--bingo", required=True, help="export agregado por conta - BINGO")
    ap.add_argument("--reals", required=True, help="export agregado por conta - REALS")
    ap.add_argument("--inicio", required=True, help="segunda-feira (AAAA-MM-DD)")
    ap.add_argument("--fim", required=True, help="domingo (AAAA-MM-DD)")
    ap.add_argument("--arquivo", default=None, help="xlsx acumulado existente (opcional)")
    ap.add_argument("--out", default=".")
    a = ap.parse_args()

    inicio, fim = pd.Timestamp(a.inicio), pd.Timestamp(a.fim)
    if inicio.dayofweek != 0 or fim.dayofweek != 6 or (fim - inicio).days != 6:
        print(f"AVISO: semana esperada SEG->DOM; recebi {inicio:%a %d/%m} -> {fim:%a %d/%m}.")
    periodo = f"{inicio:%d/%m} a {fim:%d/%m}"
    sem = _semana_label(inicio, fim)

    dados_b = {k: v for k, v in read_contas(a.bingo).items() if k not in EXCLUIDOS}
    dados_r = {k: v for k, v in read_contas(a.reals).items() if k not in EXCLUIDOS}

    if a.arquivo and Path(a.arquivo).exists():
        wb = load_workbook(a.arquivo)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    tb = _monta_aba(wb, f"B {sem}", [(b, u) for b, u, _ in BINGO], dados_b, periodo,
                    40, lambda _b: 0.35, zera=True)
    tr = _monta_aba(wb, f"R {sem}", REALS, dados_r, periodo, 30,
                    lambda b: 0.35 if b in REVSHARE35_REALS else 0.25, zera=False)
    assert tb == 12, f"TOTAL BINGO deveria estar na linha 12, ficou em {tb}"
    assert tr == 56, f"TOTAL REALS deveria estar na linha 56, ficou em {tr}"

    semanas = _reconstroi_indice(wb)

    out_dir = Path(a.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"FECHAMENTO_FINAL_PENDENTE_DESDE_290626_ATE_{fim:%d%m%y}.xlsx"
    wb.save(path)

    def conf(contas, dados, cpa, rs_fn, zera):
        tot = 0.0
        for item in contas:
            btag, user = item[0], item[1]
            rec = dados.get(btag) or dados.get(user.lower()) or {"ftd": 0.0, "ngr": 0.0}
            e = round(rec["ftd"] * cpa, 2)
            h = round(rec["ngr"] * rs_fn(btag), 2)
            tot += (0 if (rec["ftd"] == 0 or e + h < 0) else e + h) if zera else e + h
        return tot

    cb = conf(BINGO, dados_b, 40, lambda _b: 0.35, True)
    cr = conf(REALS, dados_r, 30, lambda b: 0.35 if b in REVSHARE35_REALS else 0.25, False)
    print(f"OK  {path}")
    print(f"    semana   : {periodo}  (abas 'B {sem}' / 'R {sem}')")
    print(f"    BINGO    : R$ {cb:,.2f}   (9 linhas, TOTAL na 12)")
    print(f"    REALS    : R$ {cr:,.2f}   (53 linhas, TOTAL na 56)")
    print(f"    semanas no arquivo: {len(semanas)} -> {', '.join(semanas)}")
    print("    -> rodar: python3 /mnt/skills/public/xlsx/scripts/recalc.py "
          f"{path} 90   (esperar total_errors: 0)")


if __name__ == "__main__":
    main()
