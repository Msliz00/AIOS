#!/usr/bin/env python3
"""Audita a PASTA DE EXPORTS SEMANAIS (insumo do fechamento pendente).

Esses arquivos nao tem data interna utilizavel: sao exports agregados por conta
(coluna Day vem com o literal do ano). A unica verdade de data e o NOME + a
regra da cadeia semanal: toda semana vai de SEGUNDA a DOMINGO e emenda na
anterior, sem buraco e sem sobreposicao.

  python3 auditar_exports.py "/caminho/da/pasta"
  python3 auditar_exports.py "/caminho/da/pasta" --aplicar   # corrige os nomes

O que ele checa:
  1. inicio e segunda-feira e fim e domingo, com 7 dias de intervalo
  2. cadeia sem buraco nem sobreposicao (por plataforma)
  3. BINGO e REALS pareados semana a semana
  4. o arquivo abre e tem as colunas que o pendente.py precisa (Affiliate/FTD/NGR)
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("faltou: pip install openpyxl")

# aceita 14:09:26, 14/09/26, 14-09-26, 14.09.2026
DATA_RE = re.compile(r"\b(\d{1,2})[:/.\-](\d{1,2})[:/.\-](\d{2,4})\b")
DIAS = ["segunda", "terca", "quarta", "quinta", "sexta", "sabado", "domingo"]


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return s.upper()


def plataforma(nome: str) -> str | None:
    n = _norm(nome)
    if "BINGO" in n:
        return "BINGO"
    if "REALS" in n or "REAL" in n:
        return "REALS"
    return None


def datas_do_nome(nome: str) -> list[dt.date]:
    out = []
    for d, m, a in DATA_RE.findall(nome):
        ano = int(a)
        ano += 2000 if ano < 100 else 0
        try:
            out.append(dt.date(ano, int(m), int(d)))
        except ValueError:
            pass
    return out


def checar_conteudo(caminho: Path) -> str:
    """Confere se o export tem o que o pendente.py precisa."""
    try:
        wb = load_workbook(caminho, data_only=True, read_only=True)
    except Exception as e:
        return f"nao abriu ({type(e).__name__})"
    ws = wb.worksheets[0]
    cab = []
    for row in ws.iter_rows(min_row=1, max_row=1):
        cab = [_norm(c.value).replace(" ", "") for c in row]
        break
    linhas = ws.max_row
    wb.close()
    faltando = []
    if not any(c.startswith("AFFILIATE") or c.startswith("BTAG") for c in cab):
        faltando.append("Affiliate/BTAG")
    if not any(c.startswith("FTD") for c in cab):
        faltando.append("FTD")
    if not any(c.startswith("NGR") for c in cab):
        faltando.append("NGR")
    if faltando:
        return "FALTA " + ", ".join(faltando)
    return f"ok ({max(linhas - 1, 0)} contas)"


def semana_valida(ini: dt.date, fim: dt.date) -> list[str]:
    erros = []
    if ini.weekday() != 0:
        erros.append(f"inicio {ini:%d/%m/%y} cai {DIAS[ini.weekday()]}, nao segunda")
    if fim.weekday() != 6:
        erros.append(f"fim {fim:%d/%m/%y} cai {DIAS[fim.weekday()]}, nao domingo")
    if (fim - ini).days != 6:
        erros.append(f"intervalo de {(fim - ini).days + 1} dias, nao 7")
    return erros


def semana_corrigida(ini: dt.date, fim: dt.date) -> tuple[dt.date, dt.date]:
    """Ancora na segunda-feira da semana do inicio."""
    base = ini - dt.timedelta(days=ini.weekday())
    return base, base + dt.timedelta(days=6)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pasta")
    ap.add_argument("--aplicar", action="store_true")
    a = ap.parse_args()

    base = Path(a.pasta).expanduser()
    if not base.is_dir():
        sys.exit(f"pasta nao encontrada: {base}")

    arquivos = sorted(p for p in base.rglob("*")
                      if p.suffix.lower() in {".xlsx", ".xlsm", ".csv"}
                      and not p.name.startswith("~$"))

    print(f"PASTA: {base}")
    print(f"ARQUIVOS: {len(arquivos)}\n")

    itens, sem_data, sem_plat = [], [], []
    for p in arquivos:
        plat = plataforma(p.stem)
        ds = datas_do_nome(p.stem)
        if plat is None:
            sem_plat.append(p)
        elif len(ds) < 2:
            sem_data.append(p)
        else:
            itens.append({"path": p, "plat": plat, "ini": ds[0], "fim": ds[-1]})

    por_plat = defaultdict(list)
    for it in itens:
        por_plat[it["plat"]].append(it)
    for v in por_plat.values():
        v.sort(key=lambda x: x["ini"])

    print("QUANTIDADE POR PLATAFORMA")
    for plat in ("BINGO", "REALS"):
        print(f"  {plat:<6} {len(por_plat[plat]):>3} semanas")
    if sem_plat:
        print(f"  (sem BINGO/REALS no nome: {len(sem_plat)})")
    if sem_data:
        print(f"  (sem duas datas no nome: {len(sem_data)})")
    print()

    renomear = []
    print("SEMANA A SEMANA")
    for plat in ("BINGO", "REALS"):
        if not por_plat[plat]:
            continue
        print(f"\n  == {plat} ==")
        anterior = None
        for it in por_plat[plat]:
            ini, fim, p = it["ini"], it["fim"], it["path"]
            # a semana real e sempre a segunda-feira da semana do inicio informado
            ci, cf = semana_corrigida(ini, fim)
            it["ci"], it["cf"] = ci, cf

            problemas = semana_valida(ini, fim)
            if anterior is not None:
                esperado = anterior + dt.timedelta(days=1)
                if ci > esperado:
                    falta = (ci - esperado).days // 7
                    problemas.append(
                        f"BURACO: {falta} semana(s) faltando antes "
                        f"({esperado:%d/%m/%y} ate {ci - dt.timedelta(days=1):%d/%m/%y})")
                elif ci < esperado:
                    problemas.append(f"SOBREPOE a semana anterior (que fecha {anterior:%d/%m/%y})")
            anterior = cf   # a cadeia segue pela semana CORRIGIDA, sem cascatear

            conteudo = checar_conteudo(p)
            marca = "OK " if not problemas and conteudo.startswith("ok") else "ERR"
            print(f"  {marca} {ini:%d/%m/%y} -> {fim:%d/%m/%y}  [{conteudo}]  {p.name}")
            for pr in problemas:
                print(f"        ! {pr}")

            if (ci, cf) != (ini, fim) and len(DATA_RE.findall(p.stem)) >= 2:
                novo = re.sub(DATA_RE,
                              lambda m, it=iter([ci, cf]): f"{next(it):%d:%m:%y}",
                              p.stem, count=2) + p.suffix
                if novo != p.name:
                    print(f"        -> {novo}")
                    renomear.append((p, p.with_name(novo)))

    semanas_b = {(i["ci"], i["cf"]) for i in por_plat["BINGO"]}
    semanas_r = {(i["ci"], i["cf"]) for i in por_plat["REALS"]}
    so_b, so_r = semanas_b - semanas_r, semanas_r - semanas_b
    print("\nPAREAMENTO BINGO x REALS")
    if not so_b and not so_r:
        print("  todas as semanas tem os dois exports")
    for ini, fim in sorted(so_b):
        print(f"  FALTA REALS de {ini:%d/%m/%y} ate {fim:%d/%m/%y}")
    for ini, fim in sorted(so_r):
        print(f"  FALTA BINGO de {ini:%d/%m/%y} ate {fim:%d/%m/%y}")

    if sem_plat or sem_data:
        print("\nNAO CLASSIFICADOS")
        for p in sem_plat:
            print(f"  ???  {p.name}  (sem BINGO/REALS no nome)")
        for p in sem_data:
            print(f"  ???  {p.name}  (nao achei duas datas no nome)")

    if not renomear:
        print("\nNENHUM NOME A CORRIGIR.")
        return
    if not a.aplicar:
        print(f"\n--- DRY-RUN. {len(renomear)} nome(s) a corrigir. "
              "Rode de novo com --aplicar. ---")
        return
    print("\nAPLICANDO")
    for p, novo in renomear:
        if novo.exists():
            print(f"  PULEI  {p.name} (destino ja existe)")
            continue
        p.rename(novo)
        print(f"  RENOMEADO  {p.name} -> {novo.name}")


if __name__ == "__main__":
    main()
