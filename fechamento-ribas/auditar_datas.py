#!/usr/bin/env python3
"""Audita uma pasta de fechamentos: confere a QUANTIDADE de material e se o
DDMMAA do NOME do arquivo bate com o ultimo dia REAL de dado dentro dele.

Por padrao so relata (dry-run). Com --aplicar, renomeia os arquivos errados.

  python3 auditar_datas.py "/caminho/da/pasta"
  python3 auditar_datas.py "/caminho/da/pasta" --aplicar

Como acha o ultimo dia real (nunca pelo titulo, que pode estar errado):
  - abas com cabecalho 'DIA'      -> maior data da coluna DIA (Jean / VT)
  - abas com cabecalho 'PERIODO'  -> maior data citada no periodo (pendente semanal)
Se um arquivo nao tiver nenhuma das duas, e reportado como NAO AUDITAVEL.
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("faltou: pip install openpyxl")

DATA_RE = re.compile(r"\b(\d{2})/(\d{2})/(\d{2,4})\b")
ATE_RE = re.compile(r"_ATE_(\d{6})(?=\.|_|$)", re.IGNORECASE)


def _norm(v) -> str:
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode()
    return s.strip().upper()


def _datas_em(valor) -> list[dt.date]:
    if isinstance(valor, dt.datetime):
        return [valor.date()]
    if isinstance(valor, dt.date):
        return [valor]
    out = []
    for d, m, a in DATA_RE.findall(str(valor or "")):
        ano = int(a)
        ano += 2000 if ano < 100 else 0
        try:
            out.append(dt.date(ano, int(m), int(d)))
        except ValueError:
            pass
    return out


def ultimo_dia_real(caminho: Path):
    """Devolve (data, origem) ou (None, motivo)."""
    try:
        wb = load_workbook(caminho, data_only=True, read_only=True)
    except Exception as e:
        return None, f"nao abriu ({type(e).__name__})"

    melhor, origem = None, None
    for ws in wb.worksheets:
        col_dia = col_per = header_row = None
        for row in ws.iter_rows(min_row=1, max_row=12):
            for cell in row:
                t = _norm(cell.value)
                if t == "DIA":
                    col_dia, header_row = cell.column, cell.row
                elif t in {"PERIODO", "PERIODO "}:
                    col_per, header_row = cell.column, cell.row
            if col_dia or col_per:
                break
        if header_row is None:
            continue

        alvo = col_dia or col_per
        rotulo = "coluna DIA" if col_dia else "coluna PERIODO"
        for row in ws.iter_rows(min_row=header_row + 1, min_col=alvo, max_col=alvo):
            for d in _datas_em(row[0].value):
                if melhor is None or d > melhor:
                    melhor, origem = d, f"{ws.title} / {rotulo}"
    wb.close()
    if melhor is None:
        return None, "sem coluna DIA nem PERIODO"
    return melhor, origem


def trilha(nome: str) -> str:
    n = nome.upper()
    if "JEAN" in n:
        return "JEAN"
    if "VT" in n:
        return "VT"
    if "PENDENTE" in n:
        return "PENDENTE"
    return "OUTROS"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pasta")
    ap.add_argument("--aplicar", action="store_true",
                    help="renomeia de fato (sem isso, so relata)")
    a = ap.parse_args()

    base = Path(a.pasta).expanduser()
    if not base.is_dir():
        sys.exit(f"pasta nao encontrada: {base}")

    arquivos = sorted(p for p in base.rglob("*")
                      if p.suffix.lower() in {".xlsx", ".xlsm"}
                      and not p.name.startswith("~$"))

    print(f"PASTA: {base}")
    print(f"ARQUIVOS .xlsx: {len(arquivos)}\n")

    por_trilha = defaultdict(list)
    for p in arquivos:
        por_trilha[trilha(p.name)].append(p)
    print("QUANTIDADE DE MATERIAL POR TRILHA")
    for t in ("JEAN", "VT", "PENDENTE", "OUTROS"):
        if por_trilha[t]:
            print(f"  {t:<9} {len(por_trilha[t]):>3}")
    print()

    ok, errados, sem_data, nao_audit = [], [], [], []
    vistos = defaultdict(list)

    for p in arquivos:
        real, origem = ultimo_dia_real(p)
        m = ATE_RE.search(p.stem)
        if real is None:
            nao_audit.append((p, origem))
            continue
        vistos[(trilha(p.name), real)].append(p.name)
        alvo = f"{real:%d%m%y}"
        if m is None:
            sem_data.append((p, real, origem))
        elif m.group(1) != alvo:
            novo = p.with_name(ATE_RE.sub(f"_ATE_{alvo}", p.name))
            errados.append((p, novo, m.group(1), real, origem))
        else:
            ok.append((p, real))

    print(f"NOME CONFERE ({len(ok)})")
    for p, real in ok:
        print(f"  OK   {p.name}   [dado ate {real:%d/%m/%Y}]")

    print(f"\nNOME ERRADO ({len(errados)})")
    for p, novo, tinha, real, origem in errados:
        print(f"  ERR  {p.name}")
        print(f"       nome diz _ATE_{tinha} / dado vai ate {real:%d/%m/%Y}  ({origem})")
        print(f"       ->   {novo.name}")

    if sem_data:
        print(f"\nSEM _ATE_ NO NOME ({len(sem_data)})")
        for p, real, origem in sem_data:
            print(f"  ...  {p.name}   [dado ate {real:%d/%m/%Y}]")

    if nao_audit:
        print(f"\nNAO AUDITAVEL ({len(nao_audit)})")
        for p, motivo in nao_audit:
            print(f"  ???  {p.name}   ({motivo})")

    dups = {k: v for k, v in vistos.items() if len(v) > 1}
    if dups:
        print(f"\nDUPLICATAS (mesma trilha + mesmo ultimo dia)")
        for (t, d), nomes in sorted(dups.items(), key=lambda x: str(x[0])):
            print(f"  {t} ate {d:%d/%m/%Y}: {len(nomes)} arquivos")
            for n in nomes:
                print(f"     - {n}")

    if not a.aplicar:
        print(f"\n--- DRY-RUN. {len(errados)} arquivo(s) a renomear. "
              "Rode de novo com --aplicar para efetivar. ---")
        return

    print("\nAPLICANDO")
    for p, novo, *_ in errados:
        if novo.exists():
            print(f"  PULEI  {p.name} -> {novo.name} (destino ja existe)")
            continue
        p.rename(novo)
        print(f"  RENOMEADO  {p.name} -> {novo.name}")
    print(f"\n{len(errados)} arquivo(s) processado(s).")


if __name__ == "__main__":
    main()
