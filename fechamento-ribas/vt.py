#!/usr/bin/env python3
"""Gerador VT UNIFICADO (38438 + 40271 + 38541) - ciclo MENSAL.

Uso:
  python3 vt.py --e38438 a.csv --e40271 b.csv --e38541 c.csv --mes 2026-09 [--ate 2026-09-17]

Formato (fixo):
  - Abas individuais: 13 colunas A->M (COM CPA / REVSHARE 35% / COMISSAO DIA), SEM 'CUSTO P/ FTD'.
  - Aba 'VT UNIFICADO' (primeira): 10 colunas, SEM comissao, puxa por formula das 3 abas.
  - Comissao do mes = soma do COMISSAO DIA (total M) das 3 abas.
"""
from __future__ import annotations

import argparse
import calendar
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from common import (CENTER, FILL_BG, INT, LEFT, MONEY, PCT, aplicar_estilo,
                    checar_ultimo_dia, f, filtrar, read_export)

BTAGS = ["38438", "40271", "38541"]
NOMES = {"38438": "Bigthiago11", "40271": "Instahacker", "38541": "Bigthiago01"}
MESES = ["JANEIRO", "FEVEREIRO", "MARCO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]

BASE10 = ["DIA", "REGISTROS", "FTD", "% REG P/ FTD", "VALOR DEP.", "NGR",
          "MÉDIA DEP/DIA", "MÉDIA NGR/DIA", "% DEP → NGR", "TICKET MÉD. DEP."]
HEAD13 = BASE10 + ["CPA", "REVSHARE 35%", "COMISSÃO DIA"]

HEADER_ROW = 3
FIRST = HEADER_ROW + 1


def _titulo(ws, texto, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=texto)
    c.font = f(size=14, bold=True, color="F4C430")
    c.alignment = CENTER


def _metricas(ws, t, ncols, com_comissao: bool):
    m = t + 2
    cm = ws.cell(row=m, column=1, value="MÉTRICAS GERAIS")
    cm.font = f(size=12, bold=True, color="F4C430")
    cm.alignment = LEFT
    linhas = [
        ("Total Dias", f"=COUNT(B{FIRST}:B{t - 1})", INT),
        ("Total Registros", f"=B{t}", INT),
        ("Total FTD", f"=C{t}", INT),
        ("Total Valor Depositado", f"=E{t}", MONEY),
        ("Total NGR", f"=F{t}", MONEY),
        ("Ticket Médio Depósito", f"=J{t}", MONEY),
        ("Ticket Médio NGR", f"=IF(C{t}=0,0,F{t}/C{t})", MONEY),
        ("Conversão Registro→FTD", f"=D{t}", PCT),
        ("% Depósito→NGR", f"=I{t}", PCT),
    ]
    if com_comissao:
        linhas.append(("COMISSÃO TOTAL (BINGO)", f"=M{t}", MONEY))
    for i, (label, val, fmt) in enumerate(linhas, start=1):
        r = m + i
        ws.cell(row=r, column=1, value=label).alignment = LEFT
        ws.cell(row=r, column=2, value=val).number_format = fmt
    return m + len(linhas)


def aba_individual(wb, btag, df, ate):
    ws = wb.create_sheet(f"BTAG {btag}")
    _titulo(ws, f"{NOMES.get(btag, btag)} ({btag}) — ATÉ {ate:%d/%m/%y}", 13)
    for i, h in enumerate(HEAD13, start=1):
        ws.cell(row=HEADER_ROW, column=i, value=h)

    for i, rec in enumerate(df.itertuples(index=False)):
        r = FIRST + i
        ws.cell(row=r, column=1, value=rec.day.strftime("%d/%m/%Y"))
        ws.cell(row=r, column=2, value=int(round(rec.registros))).number_format = INT
        ws.cell(row=r, column=3, value=int(round(rec.ftd))).number_format = INT
        ws.cell(row=r, column=4, value=f"=IF(B{r}=0,0,C{r}/B{r})").number_format = PCT
        ws.cell(row=r, column=5, value=float(rec.deposits)).number_format = MONEY
        ws.cell(row=r, column=6, value=float(rec.ngr)).number_format = MONEY
        ws.cell(row=r, column=7, value=f"=E{r}").number_format = MONEY
        ws.cell(row=r, column=8, value=f"=F{r}").number_format = MONEY
        ws.cell(row=r, column=9, value=f"=IF(E{r}=0,0,F{r}/E{r})").number_format = PCT
        ws.cell(row=r, column=10, value=f"=IF(C{r}=0,0,E{r}/C{r})").number_format = MONEY
        ws.cell(row=r, column=11, value=f"=C{r}*40").number_format = MONEY
        ws.cell(row=r, column=12, value=f"=ROUND(F{r}*0.35,2)").number_format = MONEY
        ws.cell(row=r, column=13,
                value=f"=IF(C{r}=0,0,IF(K{r}+L{r}<0,0,K{r}+L{r}))").number_format = MONEY

    last = FIRST + len(df) - 1
    t = last + 1
    ws.cell(row=t, column=1, value="2026")
    for col in ("B", "C"):
        ws.cell(row=t, column=ord(col) - 64,
                value=f"=SUM({col}{FIRST}:{col}{last})").number_format = INT
    ws.cell(row=t, column=4, value=f"=IF(B{t}=0,0,C{t}/B{t})").number_format = PCT
    ws.cell(row=t, column=5, value=f"=SUM(E{FIRST}:E{last})").number_format = MONEY
    ws.cell(row=t, column=6, value=f"=SUM(F{FIRST}:F{last})").number_format = MONEY
    ws.cell(row=t, column=7, value=f"=AVERAGE(E{FIRST}:E{last})").number_format = MONEY
    ws.cell(row=t, column=8, value=f"=AVERAGE(F{FIRST}:F{last})").number_format = MONEY
    ws.cell(row=t, column=9, value=f"=IF(E{t}=0,0,F{t}/E{t})").number_format = PCT
    ws.cell(row=t, column=10, value=f"=IF(C{t}=0,0,E{t}/C{t})").number_format = MONEY
    ws.cell(row=t, column=11, value=f"=SUM(K{FIRST}:K{last})").number_format = MONEY
    ws.cell(row=t, column=12, value=f"=SUM(L{FIRST}:L{last})").number_format = MONEY
    ws.cell(row=t, column=13, value=f"=SUM(M{FIRST}:M{last})").number_format = MONEY

    fim_bloco = _metricas(ws, t, 13, com_comissao=True)
    aplicar_estilo(ws, 13, HEADER_ROW, FIRST, t, ngr_col=6, ultima_linha=fim_bloco)
    ws.cell(row=1, column=1).fill = FILL_BG
    return t


def aba_unificada(wb, dias, totais_row, mes_label, ate):
    ws = wb.create_sheet("VT UNIFICADO", 0)
    _titulo(ws, f"VT UNIFICADO (38438 + 40271 + 38541) — {mes_label} ATÉ {ate:%d/%m/%y}", 10)
    for i, h in enumerate(BASE10, start=1):
        ws.cell(row=HEADER_ROW, column=i, value=h)

    def soma(col, r):
        return "=" + "+".join(f"'BTAG {b}'!{col}{r}" for b in BTAGS)

    for i, day in enumerate(dias):
        r = FIRST + i
        ws.cell(row=r, column=1, value=day.strftime("%d/%m/%Y"))
        ws.cell(row=r, column=2, value=soma("B", r)).number_format = INT
        ws.cell(row=r, column=3, value=soma("C", r)).number_format = INT
        ws.cell(row=r, column=4, value=f"=IF(B{r}=0,0,C{r}/B{r})").number_format = PCT
        ws.cell(row=r, column=5, value=soma("E", r)).number_format = MONEY
        ws.cell(row=r, column=6, value=soma("F", r)).number_format = MONEY
        ws.cell(row=r, column=7, value=f"=E{r}").number_format = MONEY
        ws.cell(row=r, column=8, value=f"=F{r}").number_format = MONEY
        ws.cell(row=r, column=9, value=f"=IF(E{r}=0,0,F{r}/E{r})").number_format = PCT
        ws.cell(row=r, column=10, value=f"=IF(C{r}=0,0,E{r}/C{r})").number_format = MONEY

    last = FIRST + len(dias) - 1
    t = last + 1
    ws.cell(row=t, column=1, value="2026")
    ws.cell(row=t, column=2, value=f"=SUM(B{FIRST}:B{last})").number_format = INT
    ws.cell(row=t, column=3, value=f"=SUM(C{FIRST}:C{last})").number_format = INT
    ws.cell(row=t, column=4, value=f"=IF(B{t}=0,0,C{t}/B{t})").number_format = PCT
    ws.cell(row=t, column=5, value=f"=SUM(E{FIRST}:E{last})").number_format = MONEY
    ws.cell(row=t, column=6, value=f"=SUM(F{FIRST}:F{last})").number_format = MONEY
    ws.cell(row=t, column=7, value=f"=AVERAGE(E{FIRST}:E{last})").number_format = MONEY
    ws.cell(row=t, column=8, value=f"=AVERAGE(F{FIRST}:F{last})").number_format = MONEY
    ws.cell(row=t, column=9, value=f"=IF(E{t}=0,0,F{t}/E{t})").number_format = PCT
    ws.cell(row=t, column=10, value=f"=IF(C{t}=0,0,E{t}/C{t})").number_format = MONEY

    fim = _metricas(ws, t, 10, com_comissao=False)
    # comissao consolidada fica fora do bloco de metricas da unificada (referencia gerencial)
    r = fim + 2
    ws.cell(row=r, column=1, value="COMISSÃO VT DO MÊS (soma das 3 abas)").alignment = LEFT
    ws.cell(row=r, column=1).font = f(bold=True, color="F4C430")
    ws.cell(row=r, column=2,
            value="=" + "+".join(f"'BTAG {b}'!M{totais_row[b]}" for b in BTAGS)
            ).number_format = MONEY
    aplicar_estilo(ws, 10, HEADER_ROW, FIRST, t, ngr_col=6, ultima_linha=r)
    ws.cell(row=1, column=1).fill = FILL_BG


def main():
    ap = argparse.ArgumentParser()
    for b in BTAGS:
        ap.add_argument(f"--e{b}", required=True, help=f"export diario BTAG {b}")
    ap.add_argument("--mes", required=True, help="AAAA-MM (ex.: 2026-09)")
    ap.add_argument("--ate", default=None)
    ap.add_argument("--out", default=".")
    a = ap.parse_args()

    ano, mes = (int(x) for x in a.mes.split("-"))
    inicio = pd.Timestamp(ano, mes, 1)
    fim_mes = pd.Timestamp(ano, mes, calendar.monthrange(ano, mes)[1])

    brutos = {b: read_export(getattr(a, f"e{b}")) for b in BTAGS}
    ate = min(checar_ultimo_dia(df, a.ate) for df in brutos.values())
    ate = min(ate, fim_mes)

    dfs = {b: filtrar(df, inicio, ate) for b, df in brutos.items()}
    dias = sorted({d for df in dfs.values() for d in df["day"]})
    if not dias:
        raise SystemExit(f"Nenhum dado em {a.mes} ate {ate:%d/%m/%Y}.")

    # alinha as 3 abas nas mesmas linhas/datas (formulas da unificada dependem disso)
    grade = pd.DataFrame({"day": dias})
    for b in BTAGS:
        dfs[b] = (grade.merge(dfs[b], on="day", how="left")
                  .fillna({"registros": 0, "ftd": 0, "deposits": 0.0, "ngr": 0.0}))

    wb = Workbook()
    wb.remove(wb.active)
    totais_row = {b: aba_individual(wb, b, dfs[b], ate) for b in BTAGS}
    mes_label = MESES[mes - 1]
    aba_unificada(wb, dias, totais_row, mes_label, ate)

    out_dir = Path(a.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"FECHAMENTO_VT_UNIFICADO_{mes_label}_ATE_{ate:%d%m%y}.xlsx"
    wb.save(path)

    comissao = 0.0
    for b in BTAGS:
        d = dfs[b]
        dia = (d["ftd"] * 40 + (d["ngr"] * 0.35).round(2))
        comissao += float(dia.where((d["ftd"] > 0) & (dia >= 0), 0).sum())

    print(f"OK  {path}")
    print(f"    periodo  : {inicio:%d/%m/%Y} -> {ate:%d/%m/%Y}  ({len(dias)} dias)")
    for b in BTAGS:
        d = dfs[b]
        print(f"    {b:>6}  ftd={d.ftd.sum():>6.0f}  ngr=R$ {d.ngr.sum():>14,.2f}")
    print(f"    COMISSAO VT (conferencia): R$ {comissao:,.2f}")
    print("    -> rodar: python3 /mnt/skills/public/xlsx/scripts/recalc.py "
          f"{path} 90   (esperar total_errors: 0)")


if __name__ == "__main__":
    main()
