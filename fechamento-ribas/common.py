"""Utilidades compartilhadas dos fechamentos Ribas/Thiagao.

Regras fixas (ver README.md):
  BINGO  -> CPA R$40/FTD + 35% RevShare, comissao do dia zera se FTD=0 ou soma<0
  REALS  -> CPA R$30/FTD + RevShare 35% (3 contas Daniel) ou 25% (demais), mantem negativo
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- visual Play
BG = "000000"
GOLD = "F4C430"
GREEN = "4ADE80"
RED = "F87171"
WHITE = "FFFFFF"
FONT_NAME = "Play"
MONEY = '"R$ "#,##0.00'
PCT = "0.00%"
INT = "#,##0"

FILL_BG = PatternFill("solid", fgColor=BG)
FILL_GOLD = PatternFill("solid", fgColor=GOLD)


def f(size=11, bold=False, color=WHITE):
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


# ------------------------------------------------------------------- parsing
def parse_br(value) -> float:
    """Converte numero em formato BR ('1.234,56', 'R$ 1.234,56', '(123,45)') para float."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return 0.0 if pd.isna(value) else float(value)
    s = str(value).strip()
    if not s or s.lower() in {"nan", "-", "--", "none"}:
        return 0.0
    neg = s.startswith("(") and s.endswith(")")
    s = re.sub(r"[^\d,.\-]", "", s)
    if not s or s in {"-", ".", ","}:
        return 0.0
    if "," in s and "." in s:
        # formato BR: ponto = milhar, virgula = decimal
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        out = float(s)
    except ValueError:
        return 0.0
    return -out if neg else out


def _norm(name: str) -> str:
    s = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


ALIASES = {
    "day": ["day", "date", "data", "dia", "reportdate", "daydate"],
    "registros": ["registrations", "registration", "registros", "registro", "signups",
                  "signup", "players", "newplayers", "cadastros"],
    "ftd": ["ftd", "ftds", "ftdcount", "firstdeposits", "firstdeposit", "firsttimedepositors",
            "qtdftd", "firstdepositors"],
    "deposits": ["deposits", "deposit", "depositamount", "depositsamount", "totaldeposits",
                 "valordepositado", "valordep", "valordeposito", "depositos"],
    "ngr": ["ngr", "netgamingrevenue", "netrevenue", "ngramount", "ngrtotal"],
}


def _pick(cols: dict, key: str) -> str | None:
    for alias in ALIASES[key]:
        if alias in cols:
            return cols[alias]
    for alias in ALIASES[key]:
        for norm, original in cols.items():
            if norm.startswith(alias) or alias in norm:
                return original
    return None


def read_export(path: str | Path, expect_affiliate: str | None = None) -> pd.DataFrame:
    """Le um export (csv/xlsx) e devolve DataFrame com day/registros/ftd/deposits/ngr.

    expect_affiliate: se informado, valida que o export e do tipo Affiliate (armadilha #2).
    """
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        raw = pd.read_excel(path, dtype=str)
    else:
        raw = None
        for sep in [None, ";", ",", "\t"]:
            try:
                raw = pd.read_csv(path, dtype=str, sep=sep, engine="python")
                if raw.shape[1] > 1:
                    break
            except Exception:
                continue
        if raw is None:
            raise SystemExit(f"Nao consegui ler {path}")

    cols = {_norm(c): c for c in raw.columns}

    if expect_affiliate is not None:
        has_aff = any(n.startswith("affiliate") for n in cols)
        has_camp = any(n.startswith("campaign") for n in cols)
        if not has_aff and has_camp:
            raise SystemExit(
                f"RECUSADO: {path.name} e export Campaign (UX-...). "
                f"Jean precisa do export Affiliate (Affiliate={expect_affiliate}). "
                "Atribuicao diferente -> totais divergem."
            )
        if has_aff:
            aff_col = next(cols[n] for n in cols if n.startswith("affiliate"))
            vals = {re.sub(r"\D", "", str(v)) for v in raw[aff_col].dropna().unique()}
            vals.discard("")
            if vals and expect_affiliate not in vals:
                raise SystemExit(
                    f"RECUSADO: {path.name} tem Affiliate={sorted(vals)}, esperado {expect_affiliate}."
                )

    day_col = _pick(cols, "day")
    if day_col is None:
        raise SystemExit(f"{path.name}: coluna de data nao encontrada. Colunas: {list(raw.columns)}")

    out = pd.DataFrame()
    out["day"] = pd.to_datetime(raw[day_col], errors="coerce", dayfirst=False)
    if out["day"].isna().all():
        out["day"] = pd.to_datetime(raw[day_col], errors="coerce", dayfirst=True)

    for key in ("registros", "ftd", "deposits", "ngr"):
        col = _pick(cols, key)
        if col is None:
            print(f"AVISO {path.name}: coluna '{key}' nao encontrada -> assumindo 0")
            out[key] = 0.0
        else:
            out[key] = raw[col].map(parse_br)

    out = out.dropna(subset=["day"])
    out = out.groupby("day", as_index=False).sum(numeric_only=True)
    out = out.sort_values("day").reset_index(drop=True)
    if out.empty:
        raise SystemExit(f"{path.name}: nenhuma linha com data valida.")
    return out


def filtrar(df: pd.DataFrame, inicio=None, fim=None) -> pd.DataFrame:
    if inicio is not None:
        df = df[df["day"] >= pd.Timestamp(inicio)]
    if fim is not None:
        df = df[df["day"] <= pd.Timestamp(fim)]
    return df.sort_values("day").reset_index(drop=True)


def checar_ultimo_dia(df: pd.DataFrame, pedido=None) -> pd.Timestamp:
    """Armadilha #1: data no nome do arquivo != ultimo dia com dado."""
    real = df["day"].max()
    if pedido is not None:
        pedido = pd.Timestamp(pedido)
        if real < pedido:
            print(f"AVISO: pedido ate {pedido:%d/%m/%Y}, mas o export so tem dados ate "
                  f"{real:%d/%m/%Y}. Fechando em {real:%d/%m/%Y}.")
        elif real > pedido:
            print(f"AVISO: export vai ate {real:%d/%m/%Y}; cortando em {pedido:%d/%m/%Y}.")
            return pedido
    return real


# ------------------------------------------------------------------ estilo
def aplicar_estilo(ws, n_cols: int, header_row: int, first_row: int,
                   total_row: int, ngr_col: int | None = None,
                   ultima_linha: int | None = None):
    ultima_linha = ultima_linha or total_row
    for row in ws.iter_rows(min_row=1, max_row=ultima_linha + 14, min_col=1, max_col=n_cols):
        for cell in row:
            cell.fill = FILL_BG
            cell.font = f()
            cell.alignment = CENTER

    for c in range(1, n_cols + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.fill = FILL_GOLD
        cell.font = f(bold=True, color="000000")
        cell = ws.cell(row=total_row, column=c)
        cell.fill = FILL_GOLD
        cell.font = f(bold=True, color="000000")

    if ngr_col:
        letra = get_column_letter(ngr_col)
        for r in range(first_row, total_row):
            cell = ws.cell(row=r, column=ngr_col)
            cell.font = f(color=GREEN)
        # cor condicional real: verde >=0, vermelho <0
        from openpyxl.formatting.rule import CellIsRule
        rng = f"{letra}{first_row}:{letra}{total_row - 1}"
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="lessThan", formula=["0"],
                            font=Font(name=FONT_NAME, size=11, color=RED)))
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="greaterThanOrEqual", formula=["0"],
                            font=Font(name=FONT_NAME, size=11, color=GREEN)))

    ws.sheet_view.showGridLines = False
    for c in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws.column_dimensions["A"].width = 14


# ------------------------------------------- leitura agregada por conta (semanal)
ALIASES_CONTA = {
    "btag": ["btag", "affiliateid", "affiliate", "subid", "tag", "idafiliado", "campaignid"],
    "username": ["username", "user", "login", "afiliado", "affiliatename", "name", "nome"],
}


def read_contas(path):
    """Le export agregado por conta -> dict {btag: {'ftd':..,'ngr':..,'username':..}}."""
    import re as _re
    from pathlib import Path as _P
    path = _P(path)
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        raw = pd.read_excel(path, dtype=str)
    else:
        raw = None
        for sep in [None, ";", ",", "\t"]:
            try:
                raw = pd.read_csv(path, dtype=str, sep=sep, engine="python")
                if raw.shape[1] > 1:
                    break
            except Exception:
                continue
        if raw is None:
            raise SystemExit(f"Nao consegui ler {path}")

    cols = {_norm(c): c for c in raw.columns}

    def pick(keys):
        for alias in keys:
            if alias in cols:
                return cols[alias]
        for alias in keys:
            for norm, original in cols.items():
                if alias in norm:
                    return original
        return None

    btag_col = pick(ALIASES_CONTA["btag"])
    user_col = pick(ALIASES_CONTA["username"])
    ftd_col = _pick(cols, "ftd")
    ngr_col = _pick(cols, "ngr")
    if btag_col is None and user_col is None:
        raise SystemExit(f"{path.name}: sem coluna de BTAG/username. Colunas: {list(raw.columns)}")

    out = {}
    for _, row in raw.iterrows():
        btag = _re.sub(r"\D", "", str(row[btag_col])) if btag_col else ""
        user = str(row[user_col]).strip() if user_col else ""
        if not btag and not user:
            continue
        key = btag or user.lower()
        rec = out.setdefault(key, {"ftd": 0.0, "ngr": 0.0, "username": user})
        rec["ftd"] += parse_br(row[ftd_col]) if ftd_col else 0.0
        rec["ngr"] += parse_br(row[ngr_col]) if ngr_col else 0.0
        if user and not rec["username"]:
            rec["username"] = user
    return out
